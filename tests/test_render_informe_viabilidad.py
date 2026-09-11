"""El generador del informe de viabilidad marca las 88 filas, no solo las que recibe.

`render_informe.py` recorría `d["preguntas"]`: las preguntas que el JSON no traía
se quedaban **sin** la marca «¿PENDIENTE ENTREVISTA?» de la columna M, que es lo que
filtra el guion de entrevista. Medido dos veces el 2026-09-10: **51 de 88** filas
marcadas en W-048UOL y **70 de 88** en W-02YZO4 (`MEJORAS`, fila #28 de `PLAN.md`,
pieza P5). Un informe con 37 filas en blanco no se lee como «37 pendientes»: se lee
como «sin cuestionario», que es lo contrario.

Los tests corren contra la **plantilla real** de `assets/` —el defecto vivía en la
relación entre el JSON y esa plantilla, así que un fixture sintético lo rodearía— y
escriben siempre en `tmp_path`, nunca en el árbol de producción.
"""
import json
import sys
from importlib import import_module
from pathlib import Path

import openpyxl
import pytest

RAIZ = Path(__file__).parent.parent
SCRIPTS = RAIZ / ".claude" / "skills" / "viabilidad-prerelleno" / "scripts"
PLANTILLA = RAIZ / ".claude" / "skills" / "viabilidad-prerelleno" / "assets" / "plantilla_informe_viabilidad.xlsx"

sys.path.insert(0, str(SCRIPTS))
render_informe = import_module("render_informe")

COL_RESPUESTA = 9    # I
COL_PENDIENTE = 13   # M

# Tres ids reales de la plantilla, del primer bloque del cuestionario.
TRES_IDS = ["cap_01", "cap_02", "cap_03"]


def _generar(tmp_path, datos: dict) -> Path:
    """Corre el generador como lo corre la skill y devuelve el .xlsx de salida."""
    entrada = tmp_path / "datos.json"
    entrada.write_text(json.dumps(datos, ensure_ascii=False), encoding="utf-8")
    salida = tmp_path / "Informe viabilidad LLM - W-TEST00.xlsx"
    argv = ["render_informe.py", str(entrada), "--salida", str(salida)]
    orig, sys.argv = sys.argv, argv
    try:
        render_informe.main()
    finally:
        sys.argv = orig
    return salida


def _mapa_de_la_plantilla() -> dict[str, int]:
    wb = openpyxl.load_workbook(PLANTILLA, read_only=True)
    try:
        return render_informe.build_id_row_map(wb["PREGUNTAS"])
    finally:
        wb.close()


def test_la_plantilla_declara_88_preguntas():
    """El «88» del que habla la pieza sale de la plantilla, no de la memoria.

    Si la plantilla cambia de cuestionario, este test lo dice en vez de dejar que
    los demás se ajusten solos al número nuevo.
    """
    assert len(_mapa_de_la_plantilla()) == 88


@pytest.fixture
def informe_con_tres_respuestas(tmp_path):
    datos = {
        "case_id": "W-TEST00",
        "preguntas": {qid: {"respuesta": f"respuesta documental {i}",
                            "cita": "[doc: prueba.pdf] cita",
                            "confianza": "alta"}
                      for i, qid in enumerate(TRES_IDS, 1)},
    }
    salida = _generar(tmp_path, datos)
    wb = openpyxl.load_workbook(salida)
    yield wb["PREGUNTAS"]
    wb.close()


def test_las_88_filas_salen_marcadas(informe_con_tres_respuestas):
    """La propiedad: ninguna pregunta del cuestionario queda sin marca.

    ROJO antes del arreglo: salían 3 de 88, que es el mismo modo de fallo que
    el 51/88 y el 70/88 medidos en producción.
    """
    ws = informe_con_tres_respuestas
    id_row = render_informe.build_id_row_map(ws)
    sin_marcar = [qid for qid, r in id_row.items()
                  if ws.cell(r, COL_PENDIENTE).value in (None, "")]
    assert not sin_marcar, (
        f"{len(sin_marcar)} de {len(id_row)} preguntas sin «¿PENDIENTE ENTREVISTA?»: "
        f"{sorted(sin_marcar)[:8]}…")


def test_las_no_respondidas_van_al_guion_y_las_respondidas_no(informe_con_tres_respuestas):
    """`sí` es «al guion de entrevista»; `no` es «resuelta en documental»."""
    ws = informe_con_tres_respuestas
    id_row = render_informe.build_id_row_map(ws)
    marcas = {qid: ws.cell(r, COL_PENDIENTE).value for qid, r in id_row.items()}

    assert [marcas[q] for q in TRES_IDS] == ["no", "no", "no"]
    resto = {q: v for q, v in marcas.items() if q not in TRES_IDS}
    assert len(resto) == 85
    assert set(resto.values()) == {"sí"}


def test_no_inventa_respuestas_al_marcar(informe_con_tres_respuestas):
    """Marcar la pendiente NO es rellenar la respuesta.

    La columna I de las 85 no resueltas tiene que seguir vacía: el valor de la
    hoja es distinguir «pendiente» de «contestada», y escribir un placeholder en I
    borraría esa distinción.
    """
    ws = informe_con_tres_respuestas
    id_row = render_informe.build_id_row_map(ws)
    con_respuesta = [qid for qid, r in id_row.items()
                     if ws.cell(r, COL_RESPUESTA).value not in (None, "")]
    assert sorted(con_respuesta) == sorted(TRES_IDS)


def test_las_filas_de_seccion_no_se_manchan(informe_con_tres_respuestas):
    """La frontera: se marcan PREGUNTAS, no filas.

    La hoja intercala cabeceras de sección («1. Captación (encargo)») que llevan
    texto en B y la columna C vacía. `build_id_row_map` ya las deja fuera; este test
    fija esa frontera para que el barrido nuevo no la cruce.
    """
    ws = informe_con_tres_respuestas
    filas_pregunta = set(render_informe.build_id_row_map(ws).values())
    manchadas = [r for r in range(5, ws.max_row + 1)
                 if r not in filas_pregunta and ws.cell(r, COL_PENDIENTE).value not in (None, "")]
    assert not manchadas, f"filas sin ID con marca escrita: {manchadas}"


def test_pendiente_explicito_del_json_manda(tmp_path):
    """El JSON puede decir `pendiente: 'no'` sin respuesta, y eso gana.

    Es el caso de una pregunta resuelta por una vía que no deja respuesta en I
    (por ejemplo, «no aplica a este tipo de caso»). El barrido nuevo no puede
    pisar una decisión explícita.
    """
    datos = {
        "case_id": "W-TEST00",
        "preguntas": {"cap_01": {"pendiente": "no"},
                      "cap_02": {"respuesta": "sí", "pendiente": "sí"}},
    }
    salida = _generar(tmp_path, datos)
    wb = openpyxl.load_workbook(salida)
    try:
        ws = wb["PREGUNTAS"]
        id_row = render_informe.build_id_row_map(ws)
        assert ws.cell(id_row["cap_01"], COL_PENDIENTE).value == "no"
        assert ws.cell(id_row["cap_02"], COL_PENDIENTE).value == "sí"
    finally:
        wb.close()


def test_una_pregunta_ajena_al_cuestionario_sigue_avisando(tmp_path, capsys):
    """El aviso de `qid` desconocido no se pierde con el barrido nuevo."""
    datos = {"case_id": "W-TEST00",
             "preguntas": {"no_existe_99": {"respuesta": "x"}}}
    _generar(tmp_path, datos)
    assert "no_existe_99" in capsys.readouterr().err


def test_las_marcas_caben_en_la_validacion_de_la_columna(informe_con_tres_respuestas):
    """`sí`/`no` son el dominio que la propia plantilla declara para M6:M103.

    La hoja PREGUNTAS trae una validación de lista `"sí,no"` sobre esa columna. Un
    literal fuera del dominio (`SÍ`, `pendiente`, `True`) no rompe nada al escribir
    —openpyxl no valida— y luego Excel se lo come al abrirlo: el defecto se vería en
    la mesa del abogado, no aquí. Este test ata el generador al dominio declarado.
    """
    ws = informe_con_tres_respuestas
    dv = next(d for d in ws.data_validations.dataValidation
              if "M6" in str(d.sqref) or "M" in str(d.sqref))
    dominio = set(dv.formula1.strip('"').split(","))
    id_row = render_informe.build_id_row_map(ws)
    escritos = {ws.cell(r, COL_PENDIENTE).value for r in id_row.values()}
    assert escritos <= dominio, f"valores fuera de la validación {dominio}: {escritos - dominio}"
